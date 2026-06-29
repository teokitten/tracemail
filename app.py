from flask import Flask, render_template, request, jsonify, send_file
import io
import os
import imaplib
import email
from email import policy
from email.parser import BytesParser
from email.utils import parsedate_to_datetime
from datetime import datetime, timezone
import PyPDF2
import docx
import reportlab
import re

app = Flask(__name__)


def _parse_folder_name(raw):
    """Extract folder name from an IMAP LIST response line."""
    if isinstance(raw, bytes):
        raw = raw.decode('utf-8', errors='replace')
    m = re.match(r'\(.*?\)\s+(?:"[^"]*"|NIL)\s+(.+)', raw.strip())
    if m:
        return m.group(1).strip().strip('"')
    return raw.rsplit(None, 1)[-1].strip('"')


@app.route('/api/connect', methods=['POST'])
def api_connect():
    data     = request.get_json(force=True)
    host     = data.get('imap_host', '').strip()
    port     = int(data.get('imap_port', 993))
    usr      = data.get('email', '').strip()
    password = data.get('password', '')

    if not host or not usr or not password:
        return jsonify({
            'ok': False,
            'error': 'Missing host, email, or password.',
            'folders': []
        })

    try:
        mail = imaplib.IMAP4_SSL(host, port)
        mail.login(usr, password)

        status, raw_list = mail.list()
        folders = []
        if status == 'OK':
            for item in raw_list:
                name = _parse_folder_name(item)
                if name:
                    folders.append(name)

        mail.logout()
        return jsonify({'ok': True, 'folders': folders, 'error': None})

    except imaplib.IMAP4.error as e:
        return jsonify({'ok': False, 'error': str(e), 'folders': []})
    except OSError as e:
        return jsonify({
            'ok': False,
            'error': f'Cannot reach {host}:{port} – {e}',
            'folders': []
        })
    except Exception as e:
        return jsonify({'ok': False, 'error': str(e), 'folders': []})


@app.route('/api/fetch', methods=['POST'])
def api_fetch():
    data        = request.get_json(force=True)
    host        = data.get('imap_host', '').strip()
    port        = int(data.get('imap_port', 993))
    usr         = data.get('email', '').strip()
    password    = data.get('password', '')
    folders     = data.get('folders', ['INBOX'])
    since_month = int(data.get('since_month', 0))
    since_year  = int(data.get('since_year', datetime.now().year))

    MONTH_ABBR = ['Jan','Feb','Mar','Apr','May','Jun',
                  'Jul','Aug','Sep','Oct','Nov','Dec']
    since_date = f'01-{MONTH_ABBR[since_month]}-{since_year}'

    try:
        mail = imaplib.IMAP4_SSL(host, port)
        mail.login(usr, password)

        seen = {}  # key: (folder, uid) → email dict

        for folder in folders:
            try:
                status, _ = mail.select(f'"{folder}"', readonly=True)
                if status != 'OK':
                    continue

                status, uid_data = mail.uid('search', None,
                                            f'SINCE {since_date}')
                if status != 'OK' or not uid_data or not uid_data[0]:
                    continue

                uid_list = uid_data[0].split()
                if not uid_list:
                    continue

                uid_str = b','.join(uid_list).decode('ascii')

                status, fetch_data = mail.uid(
                    'fetch', uid_str,
                    '(BODY.PEEK[HEADER.FIELDS (FROM TO DATE SUBJECT)])'
                )
                if status != 'OK' or not fetch_data:
                    continue

                for item in fetch_data:
                    if not isinstance(item, tuple) or len(item) < 2:
                        continue

                    fetch_info = item[0]
                    if isinstance(fetch_info, bytes):
                        fetch_info = fetch_info.decode('utf-8', errors='replace')

                    uid_match = re.search(r'UID (\d+)', fetch_info)
                    if not uid_match:
                        continue
                    uid = uid_match.group(1)

                    msg     = BytesParser(policy=policy.default).parsebytes(item[1])
                    subject  = str(msg.get('subject', '(no subject)')).strip()
                    sender   = str(msg.get('from',    '')).strip()
                    date     = str(msg.get('date',    '')).strip()
                    to_field = str(msg.get('to',      '')).strip()

                    key = (folder, uid)
                    if key not in seen:
                        seen[key] = {
                            'uid':             uid,
                            'subject':         subject,
                            'sender':          sender,
                            'to':              to_field,
                            'date':            date,
                            'has_attachments': False,
                            'snippet':         '',
                            'source_folder':   folder
                        }

            except Exception:
                continue

        mail.logout()

        def sort_key(e):
            try:
                return parsedate_to_datetime(e['date'])
            except Exception:
                return datetime.min.replace(tzinfo=timezone.utc)

        emails = sorted(seen.values(), key=sort_key, reverse=True)
        return jsonify({'emails': emails, 'count': len(emails)})

    except imaplib.IMAP4.error as e:
        return jsonify({'emails': [], 'count': 0, 'error': str(e)})
    except OSError as e:
        return jsonify({
            'emails': [], 'count': 0,
            'error': f'Cannot reach {host}:{port} – {e}'
        })
    except Exception as e:
        return jsonify({'emails': [], 'count': 0, 'error': str(e)})


@app.route('/api/thread', methods=['POST'])
def api_thread():
    data     = request.get_json(force=True)
    host     = data.get('imap_host', '').strip()
    port     = int(data.get('imap_port', 993))
    usr      = data.get('email', '').strip()
    password = data.get('password', '')
    uid      = str(data.get('uid', '')).strip()
    folder   = data.get('folder', 'INBOX')

    if not uid:
        return jsonify({'error': 'uid required'})

    try:
        mail = imaplib.IMAP4_SSL(host, port)
        mail.login(usr, password)

        status, _ = mail.select(f'"{folder}"', readonly=True)
        if status != 'OK':
            mail.logout()
            return jsonify({'error': f'Cannot select folder: {folder}'})

        status, fetch_data = mail.uid('fetch', uid, '(BODY.PEEK[])')
        mail.logout()

        if status != 'OK' or not fetch_data:
            return jsonify({'error': 'Fetch failed'})

        raw_bytes = None
        for item in fetch_data:
            if isinstance(item, tuple) and len(item) >= 2:
                raw_bytes = item[1]
                break

        if not raw_bytes:
            return jsonify({'error': 'Empty fetch response'})

        msg = BytesParser(policy=policy.default).parsebytes(raw_bytes)

        subject     = str(msg.get('subject',     '(no subject)')).strip()
        sender      = str(msg.get('from',        '')).strip()
        date        = str(msg.get('date',        '')).strip()
        in_reply_to = str(msg.get('in-reply-to', '')).strip()
        message_id  = str(msg.get('message-id',  '')).strip()

        # Body – prefer plain text, fall back to HTML
        body = ''
        if msg.is_multipart():
            for part in msg.walk():
                ct = part.get_content_type()
                cd = str(part.get('Content-Disposition') or '')
                if ct == 'text/plain' and 'attachment' not in cd:
                    raw = part.get_payload(decode=True) or b''
                    body = raw.decode(
                        part.get_content_charset() or 'utf-8',
                        errors='replace')
                    break
            if not body:
                for part in msg.walk():
                    ct = part.get_content_type()
                    cd = str(part.get('Content-Disposition') or '')
                    if ct == 'text/html' and 'attachment' not in cd:
                        raw = part.get_payload(decode=True) or b''
                        body = raw.decode(
                            part.get_content_charset() or 'utf-8',
                            errors='replace')
                        break
        else:
            raw = msg.get_payload(decode=True) or b''
            body = raw.decode(
                msg.get_content_charset() or 'utf-8',
                errors='replace')

        # Attachment metadata
        attachments = []
        walk_idx    = 0
        IMAGE_TYPES = {
            'image/jpeg', 'image/png', 'image/gif',
            'image/webp', 'image/bmp', 'image/tiff',
            'image/svg+xml'
        }
        if msg.is_multipart():
            for part in msg.walk():
                filename = part.get_filename()
                cd       = str(
                    part.get('Content-Disposition') or '').lower()
                cid      = str(
                    part.get('Content-ID') or '').strip()
                ct       = part.get_content_type().lower()

                is_image = ct in IMAGE_TYPES

                if is_image:
                    # Block all PNG – invariably signature
                    # images or logos in practice
                    if ct == 'image/png' or (
                            filename and
                            filename.lower().endswith('.png')):
                        walk_idx += 1
                        continue
                    import re as _re
                    AUTO_PAT = _re.compile(
                        r'^image\d+\.(png|jpg|jpeg|gif|bmp|webp)$',
                        _re.IGNORECASE)
                    # No filename = unidentified, block
                    if not filename:
                        walk_idx += 1
                        continue
                    # Auto-generated name = logo/signature, block
                    if AUTO_PAT.match(filename):
                        walk_idx += 1
                        continue
                    # Descriptive filename = real attachment, allow

                # Non-image parts: need filename or
                # explicit attachment disposition
                if not is_image:
                    is_explicit = cd.startswith('attachment')
                    if not is_explicit and not filename:
                        walk_idx += 1
                        continue
                    if 'inline' in cd and not filename:
                        walk_idx += 1
                        continue

                raw = part.get_payload(decode=True) or b''
                attachments.append({
                    'index':        walk_idx,
                    'name':         filename or f'attachment_{walk_idx}',
                    'content_type': ct,
                    'size_kb':      round(len(raw) / 1024, 1)
                })
                walk_idx += 1

        return jsonify({
            'subject':     subject,
            'sender':      sender,
            'date':        date,
            'body':        body,
            'in_reply_to': in_reply_to,
            'message_id':  message_id,
            'attachments': attachments
        })

    except imaplib.IMAP4.error as e:
        return jsonify({'error': str(e)})
    except OSError as e:
        return jsonify({'error': f'Cannot reach {host}:{port} – {e}'})
    except Exception as e:
        return jsonify({'error': str(e)})


@app.route('/api/attachment', methods=['POST'])
def api_attachment():
    data             = request.get_json(force=True)
    host             = data.get('imap_host', '').strip()
    port             = int(data.get('imap_port', 993))
    usr              = data.get('email', '').strip()
    password         = data.get('password', '')
    uid              = str(data.get('uid', '')).strip()
    folder           = data.get('folder', 'INBOX')
    attachment_index = int(data.get('attachment_index', 0))

    try:
        mail = imaplib.IMAP4_SSL(host, port)
        mail.login(usr, password)
        status, _ = mail.select(f'"{folder}"', readonly=True)
        if status != 'OK':
            mail.logout()
            return jsonify({'type': 'error',
                            'error': f'Cannot select {folder}'})

        status, fetch_data = mail.uid('fetch', uid, '(BODY.PEEK[])')
        mail.logout()
        if status != 'OK' or not fetch_data:
            return jsonify({'type': 'error', 'error': 'Fetch failed'})

        raw_bytes = None
        for item in fetch_data:
            if isinstance(item, tuple) and len(item) >= 2:
                raw_bytes = item[1]
                break
        if not raw_bytes:
            return jsonify({'type': 'error', 'error': 'Empty response'})

        msg = BytesParser(policy=policy.default).parsebytes(raw_bytes)

        IMAGE_TYPES = {
            'image/jpeg', 'image/png', 'image/gif',
            'image/webp', 'image/bmp', 'image/tiff',
            'image/svg+xml'
        }
        walk_idx = 0
        for part in msg.walk():
            filename = part.get_filename()
            cd       = str(
                part.get('Content-Disposition') or '').lower()
            cid      = str(
                part.get('Content-ID') or '').strip()
            ct       = part.get_content_type().lower()

            is_image = ct in IMAGE_TYPES

            if is_image:
                if ct == 'image/png' or (
                        filename and
                        filename.lower().endswith('.png')):
                    walk_idx += 1
                    continue
                import re as _re
                AUTO_PAT = _re.compile(
                    r'^image\d+\.(png|jpg|jpeg|gif|bmp|webp)$',
                    _re.IGNORECASE)
                if not filename:
                    walk_idx += 1
                    continue
                if AUTO_PAT.match(filename):
                    walk_idx += 1
                    continue
            else:
                is_explicit = cd.startswith('attachment')
                if not is_explicit and not filename:
                    walk_idx += 1
                    continue
                if 'inline' in cd and not filename:
                    walk_idx += 1
                    continue

            if walk_idx != attachment_index:
                walk_idx += 1
                continue

            name     = filename or f'attachment_{walk_idx}'
            ct       = part.get_content_type()
            raw_data = part.get_payload(decode=True) or b''

            if ct == 'application/pdf' or name.lower().endswith('.pdf'):
                import base64 as _b64
                encoded = _b64.b64encode(raw_data).decode('ascii')
                return jsonify({
                    'name': name, 'content_type': ct,
                    'type': 'pdf_binary', 'data': encoded
                })

            if (ct in (
                    'application/vnd.openxmlformats-officedocument'
                    '.wordprocessingml.document',
                    'application/msword')
                    or name.lower().endswith('.docx')
                    or name.lower().endswith('.doc')):
                try:
                    import io as _io
                    import docx as _docx
                    doc  = _docx.Document(_io.BytesIO(raw_data))
                    text = '\n'.join(
                        p.text for p in doc.paragraphs if p.text.strip())
                    return jsonify({'name': name, 'content_type': ct,
                                    'text': text or '(no text extracted)'})
                except Exception as e:
                    return jsonify({'name': name, 'content_type': ct,
                                    'text': f'DOCX extraction failed: {e}'})

            if ct.startswith('image/'):
                import base64 as _b64
                encoded = _b64.b64encode(raw_data).decode('ascii')
                return jsonify({
                    'name': name, 'content_type': ct,
                    'type': 'image', 'data': encoded
                })

            if (ct in (
                    'application/vnd.openxmlformats-officedocument'
                    '.spreadsheetml.sheet',
                    'application/vnd.ms-excel')
                    or name.lower().endswith('.xlsx')
                    or name.lower().endswith('.xls')):
                try:
                    import io as _io
                    import openpyxl as _xl
                    wb  = _xl.load_workbook(
                        _io.BytesIO(raw_data), read_only=True,
                        data_only=True)
                    sheets = []
                    for ws in wb.worksheets:
                        rows = []
                        for row in ws.iter_rows(
                                values_only=True):
                            if any(c is not None for c in row):
                                rows.append([
                                    str(c) if c is not None
                                    else ''
                                    for c in row
                                ])
                        if rows:
                            sheets.append({
                                'name': ws.title,
                                'rows': rows
                            })
                    return jsonify({
                        'name': name,
                        'content_type': ct,
                        'type': 'xlsx',
                        'sheets': sheets
                    })
                except Exception as e:
                    return jsonify({
                        'name': name, 'content_type': ct,
                        'type': 'unsupported',
                        'error': f'XLSX read failed: {e}'
                    })

            return jsonify({'type': 'unsupported', 'name': name,
                            'content_type': ct})

        return jsonify({'type': 'error',
                        'error': 'Attachment index not found'})

    except imaplib.IMAP4.error as e:
        return jsonify({'type': 'error', 'error': str(e)})
    except OSError as e:
        return jsonify({'type': 'error',
                        'error': f'Cannot reach {host}:{port} – {e}'})
    except Exception as e:
        return jsonify({'type': 'error', 'error': str(e)})


@app.route('/api/download_attachment', methods=['POST'])
def api_download_attachment():
    import io as _io
    data             = request.get_json(force=True)
    host             = data.get('imap_host', '').strip()
    port             = int(data.get('imap_port', 993))
    usr              = data.get('email', '').strip()
    password         = data.get('password', '')
    uid              = str(data.get('uid', '')).strip()
    folder           = data.get('folder', 'INBOX')
    attachment_index = int(data.get('attachment_index', 0))

    try:
        mail = imaplib.IMAP4_SSL(host, port)
        mail.login(usr, password)
        status, _ = mail.select(f'"{folder}"', readonly=True)
        if status != 'OK':
            mail.logout()
            return jsonify({'error': f'Cannot select {folder}'}), 400

        status, fetch_data = mail.uid('fetch', uid, '(BODY.PEEK[])')
        mail.logout()
        if status != 'OK' or not fetch_data:
            return jsonify({'error': 'Fetch failed'}), 500

        raw_bytes = None
        for item in fetch_data:
            if isinstance(item, tuple) and len(item) >= 2:
                raw_bytes = item[1]
                break
        if not raw_bytes:
            return jsonify({'error': 'Empty response'}), 500

        msg = BytesParser(policy=policy.default).parsebytes(raw_bytes)

        IMAGE_TYPES = {
            'image/jpeg', 'image/png', 'image/gif',
            'image/webp', 'image/bmp', 'image/tiff',
            'image/svg+xml'
        }
        walk_idx = 0
        for part in msg.walk():
            filename = part.get_filename()
            cd       = str(
                part.get('Content-Disposition') or '').lower()
            cid      = str(
                part.get('Content-ID') or '').strip()
            ct       = part.get_content_type().lower()

            is_image = ct in IMAGE_TYPES

            if is_image:
                if ct == 'image/png' or (
                        filename and
                        filename.lower().endswith('.png')):
                    walk_idx += 1
                    continue
                import re as _re
                AUTO_PAT = _re.compile(
                    r'^image\d+\.(png|jpg|jpeg|gif|bmp|webp)$',
                    _re.IGNORECASE)
                if not filename:
                    walk_idx += 1
                    continue
                if AUTO_PAT.match(filename):
                    walk_idx += 1
                    continue
            else:
                is_explicit = cd.startswith('attachment')
                if not is_explicit and not filename:
                    walk_idx += 1
                    continue
                if 'inline' in cd and not filename:
                    walk_idx += 1
                    continue

            if walk_idx != attachment_index:
                walk_idx += 1
                continue
            name     = filename or f'attachment_{walk_idx}'
            ct       = part.get_content_type()
            raw_data = part.get_payload(decode=True) or b''
            return send_file(
                _io.BytesIO(raw_data),
                mimetype=ct,
                as_attachment=True,
                download_name=name
            )

        return jsonify({'error': 'Attachment index not found'}), 404

    except imaplib.IMAP4.error as e:
        return jsonify({'error': str(e)}), 500
    except OSError as e:
        return jsonify({
            'error': f'Cannot reach {host}:{port} – {e}'
        }), 500
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route('/api/search_body', methods=['POST'])
def api_search_body():
    data          = request.get_json(force=True)
    host          = data.get('imap_host', '').strip()
    port          = int(data.get('imap_port', 993))
    usr           = data.get('email', '').strip()
    password      = data.get('password', '')
    query         = data.get('query', '').strip()
    folders       = data.get('folders', ['INBOX'])
    exclude_addrs = [
        a.strip().lower()
        for a in data.get('exclude_addrs', []) if a.strip()
    ]

    if not query:
        return jsonify({'emails': [], 'count': 0})

    try:
        mail = imaplib.IMAP4_SSL(host, port)
        mail.login(usr, password)

        seen    = {}
        q_bytes = query.encode('utf-8', errors='replace')

        for folder in folders:
            try:
                status, _ = mail.select(
                    f'"{folder}"', readonly=True)
                if status != 'OK':
                    continue

                status, uid_data = mail.uid(
                    'search', 'CHARSET', 'UTF-8',
                    'BODY', f'"{query}"')
                if status != 'OK' or not uid_data or \
                        not uid_data[0]:
                    continue

                uid_list = uid_data[0].split()
                if not uid_list:
                    continue

                uid_str = b','.join(uid_list).decode('ascii')
                status, fetch_data = mail.uid(
                    'fetch', uid_str,
                    '(BODY.PEEK[HEADER.FIELDS (FROM TO DATE SUBJECT)])'
                )
                if status != 'OK' or not fetch_data:
                    continue

                for item in fetch_data:
                    if not isinstance(item, tuple) or \
                            len(item) < 2:
                        continue
                    fetch_info = item[0]
                    if isinstance(fetch_info, bytes):
                        fetch_info = fetch_info.decode(
                            'utf-8', errors='replace')
                    uid_match = re.search(
                        r'UID (\d+)', fetch_info)
                    if not uid_match:
                        continue
                    uid = uid_match.group(1)

                    msg      = BytesParser(
                        policy=policy.default).parsebytes(item[1])
                    sender   = str(msg.get('from', '')).strip()
                    to_field = str(msg.get('to',   '')).strip()
                    date     = str(msg.get('date', '')).strip()
                    subject  = str(msg.get(
                        'subject', '(no subject)')).strip()

                    # Exclude direct correspondence
                    sender_l = sender.lower()
                    to_l     = to_field.lower()
                    skip = any(
                        addr in sender_l or addr in to_l
                        for addr in exclude_addrs
                    )
                    if skip:
                        continue

                    # Exclude automated/notification emails
                    AUTO_SENDERS = [
                        'noreply', 'no-reply', 'donotreply',
                        'do-not-reply', 'notifications@',
                        'notification@', 'alerts@', 'mailer@',
                        'postmaster@', 'bounce@', 'automated@',
                        'linkedin.com', 'facebook.com',
                        'twitter.com', 'instagram.com',
                        'youtube.com', 'google.com',
                        'greenhouse.io', 'workday.com',
                        'successfactors', 'taleo.net',
                        'jobvite.com', 'lever.co',
                        'welcometothejungle', 'xing.com'
                    ]
                    if any(s in sender_l for s in AUTO_SENDERS):
                        continue

                    # Exclude notification-style subjects
                    subject_l = subject.lower()
                    NOTIF_PATTERNS = [
                        'appeared in', 'viewed your profile',
                        'new connection', 'invitation to connect',
                        'people you may know', 'job alert',
                        'you have a new message',
                        'newsletter', 'unsubscribe',
                        '404', '500', 'error notification',
                        'delivery failed', 'out of office',
                        'automatic reply', 'automatische antwort',
                        'abwesenheitsnotiz'
                    ]
                    if any(p in subject_l for p in NOTIF_PATTERNS):
                        continue

                    key = (folder, uid)
                    if key not in seen:
                        seen[key] = {
                            'uid':            uid,
                            'subject':        subject,
                            'sender':         sender,
                            'to':             to_field,
                            'date':           date,
                            'has_attachments': False,
                            'snippet':        '',
                            'source_folder':  folder
                        }

            except imaplib.IMAP4.error:
                continue
            except Exception:
                continue

        mail.logout()

        def sort_key(e):
            try:
                return parsedate_to_datetime(e['date'])
            except Exception:
                return datetime.min.replace(tzinfo=timezone.utc)

        emails = sorted(seen.values(), key=sort_key, reverse=True)
        return jsonify({'emails': emails, 'count': len(emails)})

    except imaplib.IMAP4.error as e:
        return jsonify({'emails': [], 'count': 0, 'error': str(e)})
    except OSError as e:
        return jsonify({
            'emails': [], 'count': 0,
            'error': f'Cannot reach {host}:{port} – {e}'
        })
    except Exception as e:
        return jsonify({
            'emails': [], 'count': 0, 'error': str(e)
        })


@app.route('/api/preview', methods=['POST'])
def api_preview():
    data     = request.get_json(force=True)
    host     = data.get('imap_host', '').strip()
    port     = int(data.get('imap_port', 993))
    usr      = data.get('email', '').strip()
    password = data.get('password', '')
    uid      = str(data.get('uid', '')).strip()
    folder   = data.get('folder', 'INBOX')

    try:
        mail = imaplib.IMAP4_SSL(host, port)
        mail.login(usr, password)
        status, _ = mail.select(f'"{folder}"', readonly=True)
        if status != 'OK':
            mail.logout()
            return jsonify({'preview': ''})

        status, fetch_data = mail.uid(
            'fetch', uid, '(BODY.PEEK[])')
        mail.logout()

        if status != 'OK' or not fetch_data:
            return jsonify({'preview': ''})

        raw_bytes = None
        for item in fetch_data:
            if isinstance(item, tuple) and len(item) >= 2:
                raw_bytes = item[1]
                break
        if not raw_bytes:
            return jsonify({'preview': ''})

        msg = BytesParser(policy=policy.default).parsebytes(raw_bytes)

        body = ''
        if msg.is_multipart():
            for part in msg.walk():
                ct = part.get_content_type()
                cd = str(part.get('Content-Disposition') or '')
                if ct == 'text/plain' and 'attachment' not in cd:
                    raw = part.get_payload(decode=True) or b''
                    body = raw.decode(
                        part.get_content_charset() or 'utf-8',
                        errors='replace')
                    break
            if not body:
                for part in msg.walk():
                    ct = part.get_content_type()
                    cd = str(part.get('Content-Disposition') or '')
                    if ct == 'text/html' and 'attachment' not in cd:
                        raw = part.get_payload(decode=True) or b''
                        body = raw.decode(
                            part.get_content_charset() or 'utf-8',
                            errors='replace')
                        break
        else:
            raw = msg.get_payload(decode=True) or b''
            body = raw.decode(
                msg.get_content_charset() or 'utf-8',
                errors='replace')

        lo = body.strip().lower()
        if (lo.startswith('<html') or lo.startswith('<!doctype') or
                lo.startswith('<div') or lo.startswith('<p') or
                lo.startswith('<table')):
            body = re.sub(
                r'<style[^>]*>[\s\S]*?</style>', '',
                body, flags=re.IGNORECASE)
            body = re.sub(r'<[^>]+>', ' ', body)
            body = re.sub(r'&nbsp;', ' ', body)
            body = re.sub(r'&amp;', '&', body)
            body = re.sub(r'&lt;', '<', body)
            body = re.sub(r'&gt;', '>', body)
            body = re.sub(r'\s{3,}', '\n', body).strip()

        lines = [l.strip() for l in body.split('\n') if l.strip()]
        start = 0
        if lines:
            fl = lines[0].lower()
            if any(fl.startswith(g) for g in [
                    'dear', 'hello', 'hi ', 'good', 'hallo',
                    'sehr geehrte', 'liebe']):
                start = 1

        preview = ' '.join(lines[start:start + 3])[:220]
        return jsonify({'preview': preview})

    except Exception:
        return jsonify({'preview': ''})


@app.route('/api/export_pdf', methods=['POST'])
def api_export_pdf():
    try:
        import io as _io
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.units import mm
        from reportlab.lib.colors import HexColor
        from reportlab.platypus import (SimpleDocTemplate,
            Paragraph, Spacer, HRFlowable, KeepTogether)
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.enums import TA_LEFT, TA_RIGHT

        data        = request.get_json(force=True)
        emails      = data.get('emails', [])
        filter_used = data.get('filter_used', '')
        date_range  = data.get('date_range', {})
        user_email  = data.get('user_email', '')

        def sort_key(e):
            try:
                return parsedate_to_datetime(e.get('date', ''))
            except Exception:
                return datetime.min.replace(tzinfo=timezone.utc)

        emails = sorted(emails, key=sort_key)

        PAGE_W, PAGE_H = A4
        MARGIN         = 18 * mm
        INDENT         = 42 * mm

        C_BG      = HexColor('#0f0f14')
        C_SENT    = HexColor('#7a5200')
        C_RECV    = HexColor('#4a2a80')
        C_TEXT    = HexColor('#1a1a1a')
        C_MUTED   = HexColor('#555555')
        C_NOTE    = HexColor('#5a1a8a')
        C_RULE    = HexColor('#cccccc')
        C_DATE    = HexColor('#888888')

        def draw_page(canvas_obj, doc):
            canvas_obj.saveState()
            canvas_obj.setFont('Helvetica', 8)
            canvas_obj.setFillColor(C_MUTED)
            canvas_obj.drawString(
                MARGIN, 10 * mm,
                'Generated by Tracemail – tracemail.local')
            canvas_obj.drawRightString(
                PAGE_W - MARGIN, 10 * mm,
                f'Page {doc.page}')
            canvas_obj.restoreState()

        buf = _io.BytesIO()
        doc = SimpleDocTemplate(
            buf, pagesize=A4,
            leftMargin=MARGIN, rightMargin=MARGIN,
            topMargin=MARGIN,  bottomMargin=22 * mm)

        def ps(name, **kw):
            defaults = dict(fontName='Helvetica', fontSize=9,
                            textColor=C_TEXT, leading=14)
            defaults.update(kw)
            return ParagraphStyle(name, **defaults)

        sty_title  = ps('T', fontName='Helvetica-Bold',
                        fontSize=13, spaceAfter=3)
        sty_sub    = ps('S', textColor=C_MUTED, fontSize=8,
                        spaceAfter=10)
        sty_date_div = ps('DD', textColor=C_DATE, fontSize=8,
                          alignment=1, spaceBefore=8,
                          spaceAfter=8)

        # ── Direction styles ──────────────────────────────
        SENT_BG  = HexColor('#1a1200')
        RECV_BG  = HexColor('#0d0820')
        SENT_ACC = HexColor('#c07800')
        RECV_ACC = HexColor('#7040c0')

        sty_sent_label = ps('SL',
            fontName='Helvetica-Bold', fontSize=7,
            textColor=SENT_ACC, spaceAfter=1,
            leftIndent=INDENT)
        sty_recv_label = ps('RL',
            fontName='Helvetica-Bold', fontSize=7,
            textColor=RECV_ACC, spaceAfter=1)

        sty_sent_meta = ps('SM', fontSize=8,
            textColor=C_MUTED, leftIndent=INDENT, spaceAfter=2)
        sty_recv_meta = ps('RM', fontSize=8,
            textColor=C_MUTED, spaceAfter=2)

        sty_sent_subj = ps('SS',
            fontName='Helvetica-Bold', fontSize=10,
            leftIndent=INDENT, spaceAfter=4)
        sty_recv_subj = ps('RS',
            fontName='Helvetica-Bold', fontSize=10,
            spaceAfter=4)

        sty_sent_body = ps('SB', fontSize=9, leading=13,
            leftIndent=INDENT, spaceAfter=0)
        sty_recv_body = ps('RB', fontSize=9, leading=13,
            spaceAfter=0)

        sty_note = ps('N', fontSize=8, leading=12,
            textColor=HexColor('#9966cc'),
            leftIndent=INDENT, spaceAfter=0,
            fontName='Helvetica-Oblique')
        sty_att = ps('AT', fontSize=8,
            textColor=C_MUTED, leftIndent=INDENT, spaceAfter=0)

        def strip_quotes(text):
            lines  = text.split('\n')
            clean  = []
            skip   = False
            for ln in lines:
                s = ln.strip()
                if skip: continue
                if s.startswith('>'): continue
                if s.startswith('<!--'):
                    skip = True; continue
                if (s.lower().startswith('on ') and
                        'wrote:' in s.lower()):
                    skip = True; continue
                if s.startswith('Von:') or \
                        s.startswith('Gesendet:') or \
                        s.startswith('An:') or \
                        s.startswith('Betreff:'):
                    skip = True; continue
                if re.match(r'^-{3,}', s) and (
                        'message' in s.lower() or
                        'forwarded' in s.lower()):
                    skip = True; continue
                if s in ('________________________________',
                         '-----Original Message-----',
                         '-----Ursprüngliche Nachricht-----'):
                    skip = True; continue
                clean.append(ln)
            return '\n'.join(clean).strip() or text

        def safe(text, limit=4000):
            return (str(text)[:limit]
                    .replace('&', '&amp;')
                    .replace('<', '&lt;')
                    .replace('>', '&gt;'))

        story = []

        title_text = (f'Email thread – {filter_used}'
                      if filter_used else 'Email thread')
        story.append(Paragraph(safe(title_text), sty_title))

        date_from = date_range.get('from', '')
        date_to   = date_range.get('to', '')
        story.append(Paragraph(
            f'Date range: {safe(date_from)} to {safe(date_to)}'
            f'&nbsp;&nbsp;|&nbsp;&nbsp;'
            f'Account: {safe(user_email)}', sty_sub))
        story.append(HRFlowable(
            width='100%', thickness=1,
            color=C_RULE, spaceAfter=10))

        last_date_str = ''

        for email in emails:
            folder = (email.get('source_folder', '') or '').lower()
            sender = (email.get('sender', '') or '').lower()
            sent   = ('sent' in folder or
                      user_email.lower() in sender)

            # Date divider
            date_str = ''
            try:
                d = parsedate_to_datetime(email.get('date', ''))
                date_str = d.strftime('%B %d, %Y')
            except Exception:
                pass
            if date_str and date_str != last_date_str:
                story.append(Paragraph(
                    f'———  {date_str}  '
                    f'———',
                    ps('DD', textColor=C_DATE, fontSize=8,
                       alignment=1, spaceBefore=8, spaceAfter=8)))
                last_date_str = date_str

            l_sty = sty_sent_label if sent else sty_recv_label
            m_sty = sty_sent_meta  if sent else sty_recv_meta
            s_sty = sty_sent_subj  if sent else sty_recv_subj
            b_sty = sty_sent_body  if sent else sty_recv_body

            block = []
            label = '→ You' if sent else '← Them'
            block.append(Paragraph(label, l_sty))
            block.append(Paragraph(
                f'{safe(email.get("sender", ""))}  '
                f'&nbsp;·&nbsp;  '
                f'{safe(email.get("date", ""))}', m_sty))
            block.append(Paragraph(
                safe(email.get('subject', '(no subject)')),
                s_sty))

            body = (email.get('body', '') or '').strip()
            if body:
                body = strip_quotes(body)
                for chunk in body[:3000].split('\n\n')[:25]:
                    chunk = safe(chunk).replace('\n', '<br/>')
                    if chunk.strip():
                        block.append(Paragraph(chunk, b_sty))
                        block.append(Spacer(1, 2))

            note = (email.get('note', '') or '').strip()
            if note:
                block.append(Spacer(1, 4))
                block.append(Paragraph(
                    f'✎ {safe(note)}', sty_note))

            atts = email.get('attachments', []) or []
            if atts:
                names = ', '.join(
                    safe(a.get('name', '?'))
                    for a in atts[:10])
                block.append(Paragraph(
                    f'Attachments: {names}', sty_att))

            block.append(Spacer(1, 8))
            story.append(KeepTogether(block))

        if not emails:
            story.append(Paragraph(
                'No emails in this thread.',
                sty_recv_body))

        doc.build(story,
                  onFirstPage=draw_page,
                  onLaterPages=draw_page)
        buf.seek(0)
        return send_file(
            buf,
            mimetype='application/pdf',
            as_attachment=True,
            download_name='tracemail-export.pdf')

    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route("/")
def index():
    return render_template("index.html")


if __name__ == "__main__":
    app.run(debug=True, port=5050)
